"""Synthetic test decks. Every company, person and figure here is fictional.

Run ``python tests/fixtures/build_fixtures.py`` to regenerate the committed sample files;
the test suite builds fresh copies in a temporary directory on every run.
"""

from __future__ import annotations

import io
import sys
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.lib.pagesizes import landscape, letter
from reportlab.pdfgen import canvas

#: (title, body lines) per slide - a seed-stage healthcare SaaS company, "CareLoop Health".
SAMPLE_SLIDES: list[tuple[str, list[str]]] = [
    ("CareLoop Health", ["AI discharge coordination for community hospitals",
                         "Headquartered in Columbus, Ohio · Founded in 2022"]),
    ("Problem", ["Community hospitals lose $4.1M a year to avoidable readmissions",
                 "Discharge planning is manual and takes 6 hours of nurse time per patient"]),
    ("Solution", ["CareLoop automates discharge plans and follows up with patients for 30 days",
                  "Integrates with existing EHR workflows in under two weeks"]),
    ("Market Opportunity", ["TAM $6.2B", "SAM $1.1B", "SOM $120M",
                            "1,800 community hospitals in the United States"]),
    ("Traction", ["$1.2M ARR as of June 2025", "14 paying hospitals",
                  "Revenue growth 180% YoY", "3 pilots with regional health systems",
                  "Net revenue retention 112%"]),
    ("Business Model", ["SaaS subscription: $85K per hospital per year", "Gross margin 72%"]),
    ("Competition", ["Epic", "Cerner", "Manual processes",
                     "We are the only solution with 30-day automated follow-up"]),
    ("Team", ["Dana Reyes, CEO: former VP Operations at a regional hospital network",
              "Sam Okafor, CTO: previously led machine learning at a clinical data company"]),
    ("Financial Projections", ["Projected 2028 revenue $11M", "Monthly burn $180K"]),
    ("The Ask", ["Raising $4M Seed round on a SAFE at a $20M valuation cap",
                 "$1.5M committed", "18 months of runway"]),
    ("Use of Funds", ["Sales and marketing 45%", "Product and engineering 35%", "Operations 20%"]),
    ("Milestones", ["Reach 40 hospitals", "Launch payer analytics module", "Series A readiness in 18 months"]),
]


def make_sample_pdf(path: Path, slides: list[tuple[str, list[str]]] = SAMPLE_SLIDES) -> Path:
    width, height = landscape(letter)
    pdf = canvas.Canvas(str(path), pagesize=(width, height))
    for title, lines in slides:
        pdf.setFont("Helvetica-Bold", 30)
        pdf.drawString(60, height - 90, title)
        pdf.setFont("Helvetica", 18)
        y = height - 150
        for line in lines:
            pdf.drawString(60, y, line)
            y -= 34
        pdf.showPage()
    pdf.save()
    return path


def make_sample_pptx(path: Path, slides: list[tuple[str, list[str]]] = SAMPLE_SLIDES) -> Path:
    presentation = Presentation()
    layout = presentation.slide_layouts[1]  # title and content
    for title, lines in slides:
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = title
        body = slide.placeholders[1].text_frame
        body.text = lines[0] if lines else ""
        for line in lines[1:]:
            paragraph = body.add_paragraph()
            paragraph.text = line
            paragraph.font.size = Pt(18)
    notes_slide = presentation.slides[9].notes_slide
    notes_slide.notes_text_frame.text = "Lead investor conversations are ongoing."
    extra = presentation.slides.add_slide(presentation.slide_layouts[5])
    extra.shapes.title.text = "Unit Economics"
    table = extra.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(1.5)).table
    for row, (label, value) in enumerate([("Metric", "Value"), ("CAC", "$18K"), ("LTV", "$240K")]):
        table.cell(row, 0).text = label
        table.cell(row, 1).text = value
    presentation.save(str(path))
    return path


def make_scanned_pdf(path: Path, pages: int = 5) -> Path:
    """Image-only pages: shapes and no text layer."""
    width, height = landscape(letter)
    pdf = canvas.Canvas(str(path), pagesize=(width, height))
    for index in range(pages):
        pdf.setFillGray(0.2 + 0.1 * index)
        pdf.rect(80, 80, width - 160, height - 160, fill=1, stroke=0)
        pdf.showPage()
    pdf.save()
    return path


def make_sparse_pdf(path: Path) -> Path:
    """A deck with almost no metrics, to exercise missing-data handling."""
    return make_sample_pdf(path, [
        ("Brightpath", ["Making teams better"]),
        ("Our Vision", ["A revolutionary platform for the future of work"]),
        ("Contact", ["hello@example.com"]),
    ])


def make_corrupted_pdf(path: Path, source: Path) -> Path:
    path.write_bytes(source.read_bytes()[:400])
    return path


def make_corrupted_pptx(path: Path) -> Path:
    path.write_bytes(b"PK\x03\x04" + b"\x00garbage-not-a-zip" * 20)
    return path


def make_docx_disguised_as_pptx(path: Path) -> Path:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "<document/>")
    path.write_bytes(buffer.getvalue())
    return path


def build_all(directory: Path) -> dict[str, Path]:
    directory.mkdir(parents=True, exist_ok=True)
    sample_pdf = make_sample_pdf(directory / "careloop_deck.pdf")
    return {
        "sample_pdf": sample_pdf,
        "sample_pptx": make_sample_pptx(directory / "careloop_deck.pptx"),
        "scanned_pdf": make_scanned_pdf(directory / "scanned_deck.pdf"),
        "sparse_pdf": make_sparse_pdf(directory / "sparse_deck.pdf"),
        "corrupted_pdf": make_corrupted_pdf(directory / "corrupted_deck.pdf", sample_pdf),
        "corrupted_pptx": make_corrupted_pptx(directory / "corrupted_deck.pptx"),
        "disguised_docx": make_docx_disguised_as_pptx(directory / "word_file.pptx"),
    }


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).parent
    for name, file in build_all(target).items():
        print(f"{name}: {file}")
