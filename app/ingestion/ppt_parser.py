"""PowerPoint ingestion.

PPTX: text frames, tables, chart titles/series, grouped shapes and speaker notes via
python-pptx, one DeckPage per slide.

Legacy PPT (binary Office format): python-pptx cannot read it. When LibreOffice is installed
the file is converted to PPTX in the same private workspace; otherwise the user gets an
actionable error explaining how to convert it.
"""

from __future__ import annotations

import shutil
import subprocess
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.errors import IngestError, UnsupportedFormatError
from app.ingestion.models import DeckDocument, DeckPage
from app.utils.logging import get_logger
from app.utils.text import normalize_text

LOGGER = get_logger(__name__)

SOFFICE_TIMEOUT_S = 180


def _chart_text(chart: Any) -> list[str]:
    parts: list[str] = []
    try:
        if chart.has_title and chart.chart_title.has_text_frame:
            title = chart.chart_title.text_frame.text.strip()
            if title:
                parts.append(f"Chart: {title}")
        categories = [str(c) for c in chart.plots[0].categories] if chart.plots else []
        for plot in chart.plots:
            for series in plot.series:
                values = list(series.values)
                if categories and len(categories) == len(values):
                    pairs = ", ".join(f"{c}: {v:g}" if isinstance(v, (int, float)) else f"{c}: {v}"
                                      for c, v in zip(categories, values, strict=False))
                else:
                    pairs = ", ".join(str(v) for v in values)
                parts.append(f"{series.name or 'Series'} | {pairs}")
    except Exception:  # noqa: BLE001 - chart XML varies widely; best effort
        pass
    return parts


def _shape_content(shape: Any) -> tuple[list[str], list[str]]:
    """(text lines, table rows) for a shape, descending into groups."""
    texts: list[str] = []
    tables: list[str] = []
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for member in shape.shapes:
            sub_text, sub_tables = _shape_content(member)
            texts.extend(sub_text)
            tables.extend(sub_tables)
        return texts, tables
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [" ".join(cell.text.split()) for cell in row.cells]
            if any(cells):
                tables.append(" | ".join(c for c in cells if c))
        return texts, tables
    if getattr(shape, "has_chart", False) and shape.has_chart:
        tables.extend(_chart_text(shape.chart))
        return texts, tables
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        text = shape.text_frame.text
        if text.strip():
            texts.append(text)
    return texts, tables


def parse_pptx(path: Path, *, display_name: str | None = None) -> DeckDocument:
    name = display_name or path.name
    try:
        presentation = Presentation(str(path))
    except Exception as exc:  # python-pptx raises package-specific errors
        raise IngestError(
            f"{name} could not be opened as a PowerPoint file. It may be corrupt or password-protected; "
            "re-save it from PowerPoint and try again."
        ) from exc

    pages: list[DeckPage] = []
    for number, slide in enumerate(presentation.slides, start=1):
        texts: list[str] = []
        tables: list[str] = []
        for shape in slide.shapes:
            try:
                shape_text, shape_tables = _shape_content(shape)
            except Exception:  # noqa: BLE001 - one odd shape must not sink the slide
                continue
            texts.extend(shape_text)
            tables.extend(shape_tables)
        body = normalize_text("\n".join(texts))

        title = ""
        try:
            if slide.shapes.title is not None:
                title = slide.shapes.title.text_frame.text.strip()
        except Exception:  # noqa: BLE001
            title = ""
        if not title:
            title = next((ln.strip() for ln in body.split("\n") if ln.strip()), "")

        notes = None
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = normalize_text(slide.notes_slide.notes_text_frame.text) or None

        pages.append(DeckPage(number=number, title=title[:200] or None, text=body, notes=notes, tables=tables))

    if not pages:
        raise IngestError(f"{name} contains no slides.")
    doc = DeckDocument(filename=name, source_format="pptx", pages=pages)
    LOGGER.info("Parsed PPTX: %d slides, %d low-text slides", doc.page_count, len(doc.low_text_pages))
    return doc


def find_soffice() -> str | None:
    for candidate in ("soffice", "soffice.exe", "libreoffice"):
        found = shutil.which(candidate)
        if found:
            return found
    for path in (Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
                 Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
                 Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")):
        if path.is_file():
            return str(path)
    return None


def parse_legacy_ppt(path: Path, *, display_name: str | None = None) -> DeckDocument:
    """Convert a binary .ppt to .pptx with LibreOffice (inside ``path``'s directory), then parse it."""
    name = display_name or path.name
    soffice = find_soffice()
    if soffice is None:
        raise UnsupportedFormatError(
            f"{name} is a legacy PowerPoint 97-2003 (.ppt) file. Converting it requires LibreOffice, which "
            "is not installed on this server. Open the deck in PowerPoint and use File > Save As > "
            "PowerPoint Presentation (.pptx), or export it to PDF, then upload that file."
        )
    out_dir = path.parent / "converted"
    out_dir.mkdir(exist_ok=True)
    try:
        result = subprocess.run(  # noqa: S603 - fixed argv, resolved binary, no shell
            [soffice, "--headless", "--convert-to", "pptx", "--outdir", str(out_dir), str(path)],
            capture_output=True, timeout=SOFFICE_TIMEOUT_S, check=False,
        )
    except (subprocess.TimeoutExpired, OSError) as exc:
        raise IngestError(f"Converting {name} from .ppt failed. Save it as .pptx or PDF and re-upload.") from exc
    produced = sorted(out_dir.glob("*.pptx"))
    if result.returncode != 0 or not produced:
        raise IngestError(f"{name} could not be converted from .ppt. Save it as .pptx or PDF and re-upload.")
    doc = parse_pptx(produced[0], display_name=name)
    doc.converted_from_ppt = True
    doc.warnings.append("Converted from legacy .ppt with LibreOffice; check the extracted values carefully.")
    return doc
